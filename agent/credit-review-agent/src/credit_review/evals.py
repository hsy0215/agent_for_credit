from __future__ import annotations

import hashlib
import json
import re
from dataclasses import dataclass
from difflib import SequenceMatcher
from pathlib import Path
from typing import Any, Iterable

from openpyxl import load_workbook
from pydantic import BaseModel


class DatasetRecord(BaseModel):
    record_id: str
    company_name: str
    document_path: str
    approved_amount: float | None
    split: str
    match_score: float
    gold: dict[str, Any]


def normalize_company(value: str) -> str:
    value = re.sub(r"[\s（）()·•\-—_]", "", value)
    value = re.sub(r"(?:评审报告|签批版|以此为准|PDF版本|报告)$", "", value, flags=re.I)
    return value.replace("股份有限公司", "").replace("有限责任公司", "").replace("有限公司", "")


def company_split(company: str) -> str:
    bucket = int(hashlib.sha256(normalize_company(company).encode()).hexdigest()[:8], 16) % 10
    return "development" if bucket < 6 else "validation" if bucket < 8 else "test"


class HistoricalDatasetBuilder:
    def build(self, document_root: Path, amount_workbook: Path) -> list[DatasetRecord]:
        labels = self._read_amounts(amount_workbook)
        records: list[DatasetRecord] = []
        for path in sorted(document_root.rglob("*")):
            if path.suffix.lower() not in {".docx", ".pdf"} or "__MACOSX" in path.parts:
                continue
            candidate = normalize_company(path.stem)
            company, amount, score = self._best_match(candidate, labels)
            record_id = "DS-" + hashlib.sha256(str(path).encode()).hexdigest()[:12]
            records.append(DatasetRecord(
                record_id=record_id,
                company_name=company or path.stem,
                document_path=str(path.resolve()),
                approved_amount=amount,
                split=company_split(company or path.stem),
                match_score=score,
                gold={
                    "company_name": company or None,
                    "approved_amount": amount,
                    "established_year": None,
                    "registered_capital": None,
                    "requested_amount": None,
                    "industry": None,
                    "financials": [],
                },
            ))
        return records

    @staticmethod
    def write(records: Iterable[DatasetRecord], output: Path) -> None:
        output.parent.mkdir(parents=True, exist_ok=True)
        output.write_text(
            "\n".join(record.model_dump_json() for record in records) + "\n",
            encoding="utf-8",
        )

    @staticmethod
    def _read_amounts(path: Path) -> dict[str, tuple[str, float]]:
        workbook = load_workbook(path, read_only=True, data_only=True)
        labels: dict[str, tuple[str, float]] = {}
        for sheet in workbook:
            for row in sheet.iter_rows(min_row=2, values_only=True):
                if len(row) >= 3 and row[1] and isinstance(row[2], (int, float)):
                    company = str(row[1]).strip()
                    labels[normalize_company(company)] = (company, float(row[2]))
        return labels

    @staticmethod
    def _best_match(candidate: str, labels: dict[str, tuple[str, float]]) -> tuple[str | None, float | None, float]:
        if candidate in labels:
            company, amount = labels[candidate]
            return company, amount, 1.0
        scored = [
            (SequenceMatcher(None, candidate, normalized).ratio(), company, amount)
            for normalized, (company, amount) in labels.items()
        ]
        score, company, amount = max(scored, default=(0.0, None, None))
        return (company, amount, score) if score >= 0.72 else (None, None, score)


@dataclass
class EvalSummary:
    records: int
    dataset_records: int
    missing_predictions: int
    field_exact_accuracy: float
    field_coverage: float
    unsupported_value_rate: float
    amount_mae: float | None
    amount_within_50_accuracy: float | None
    amount_band_accuracy: float | None

    def to_dict(self) -> dict[str, Any]:
        return self.__dict__


def evaluate_jsonl(gold_path: Path, prediction_path: Path) -> EvalSummary:
    gold = {item["record_id"]: item for item in _read_jsonl(gold_path)}
    predictions = {item["record_id"]: item for item in _read_jsonl(prediction_path)}
    fields = ("company_name", "established_year", "registered_capital", "requested_amount", "industry")
    total = correct = covered = unsupported = 0
    amount_errors: list[float] = []
    within_50: list[bool] = []
    band_correct: list[bool] = []
    evaluated_ids = gold.keys() & predictions.keys()
    for record_id in evaluated_ids:
        gold_record = gold[record_id]
        predicted = predictions[record_id]
        gold_facts = gold_record.get("gold", {})
        predicted_facts = predicted.get("facts", {})
        for field in fields:
            expected = gold_facts.get(field)
            actual = _value(predicted_facts.get(field))
            if expected is None:
                if actual is not None:
                    unsupported += 1
                continue
            total += 1
            if actual is not None:
                covered += 1
            if _normalized_value(actual) == _normalized_value(expected):
                correct += 1
        expected_amount = gold_record.get("approved_amount")
        predicted_amount = predicted.get("suggested_amount")
        if expected_amount is not None and predicted_amount is not None:
            error = abs(float(predicted_amount) - float(expected_amount))
            amount_errors.append(error)
            within_50.append(error <= 50)
            band_correct.append((float(predicted_amount) <= 300) == (float(expected_amount) <= 300))
    denominator = total or 1
    attempted = total + unsupported
    return EvalSummary(
        records=len(evaluated_ids),
        dataset_records=len(gold),
        missing_predictions=len(gold) - len(evaluated_ids),
        field_exact_accuracy=correct / denominator,
        field_coverage=covered / denominator,
        unsupported_value_rate=unsupported / (attempted or 1),
        amount_mae=sum(amount_errors) / len(amount_errors) if amount_errors else None,
        amount_within_50_accuracy=sum(within_50) / len(within_50) if within_50 else None,
        amount_band_accuracy=sum(band_correct) / len(band_correct) if band_correct else None,
    )


def _read_jsonl(path: Path) -> list[dict[str, Any]]:
    return [json.loads(line) for line in path.read_text(encoding="utf-8").splitlines() if line.strip()]


def _value(value: Any) -> Any:
    return value.get("value") if isinstance(value, dict) and "value" in value else value


def _normalized_value(value: Any) -> Any:
    return normalize_company(value) if isinstance(value, str) else value
